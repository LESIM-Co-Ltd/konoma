"""special.xlsx via openpyxl: what LibreOffice cannot write.

veryHidden / hidden sheets, an Excel Table (ListObject), a chart sheet, rich text, inline strings
(`t="inlineStr"`), string formula results (`t="str"`), booleans and error cells, space-preserved
shared strings, real dates/times/durations, defined names, a formula with no cached value.

Run with any python3 that has openpyxl:  python3 gen_openpyxl.py
"""
import datetime
import os
import re
import shutil
import tempfile
import zipfile

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    rows = [["地域", "商品", "数量", "金額"], ["東", "りんご", 3, 360], ["西", "みかん", 5, 400],
            ["東", "ぶどう", 2, 900], ["西", "りんご", 7, 840], ["東", "みかん", 1, 80]]
    for r in rows:
        ws.append(r)
    tab = Table(displayName="売上表", ref="A1:D6")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)
    # a table with a totals row on the side
    ws["F1"] = "合計(式)"
    ws["F2"] = "=SUBTOTAL(109,売上表[金額])"  # no cached value: openpyxl writes none
    ws["F3"] = "=SUM(D2:D6)"

    # rich text, spaces, escapes
    rt = wb.create_sheet("リッチ")
    rt["A1"] = CellRichText("普通", TextBlock(InlineFont(b=True, color="FF0000"), "太字赤"), " と ",
                            TextBlock(InlineFont(i=True), "斜体"))
    rt["A2"] = "  leading and trailing  "
    rt["A3"] = "line1\nline2"
    rt["A4"] = "x_x000B_vt"  # the OOXML escape of a vertical tab
    rt["A5"] = "_x0041_ literal escape text"
    rt["A6"] = datetime.datetime(2026, 10, 5, 13, 30, 15)
    rt["A7"] = datetime.date(2026, 10, 5)
    rt["A8"] = datetime.time(7, 5, 9)
    rt["A9"] = datetime.timedelta(hours=36, minutes=30)
    rt["B1"] = True
    rt["B2"] = False
    rt["B3"] = "#DIV/0!"
    rt["B3"].data_type = "e"
    rt["B4"] = "#N/A"
    rt["B4"].data_type = "e"
    rt["C1"] = 1.5e300
    rt["C2"] = -0.0
    rt["C3"] = 123456789.123456789
    rt["C4"] = 10 ** 15
    rt["C5"] = 2 ** 53 + 1

    # hidden and very hidden sheets, and a chart sheet
    hid = wb.create_sheet("隠し")
    hid["A1"] = "hidden"
    hid.sheet_state = "hidden"
    vh = wb.create_sheet("超隠し")
    vh["A1"] = "veryHidden"
    vh.sheet_state = "veryHidden"
    cs = wb.create_chartsheet("グラフシート")
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=6, max_col=4), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=6))
    cs.add_chart(ch)
    last = wb.create_sheet("最後の表")
    last["A1"] = "after the chart sheet"
    wb.defined_names["金額範囲"] = DefinedName("金額範囲", attr_text="'表'!$D$2:$D$6")

    tmp = os.path.join(OUT, "special.tmp.xlsx")
    wb.save(tmp)
    # post-process: turn some cells into t="str" (string formula result), t="inlineStr"
    final = os.path.join(OUT, "special.xlsx")
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(final, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        wbxml = zin.read("xl/workbook.xml").decode()
        for n in names:
            data = zin.read(n)
            if n == "xl/worksheets/sheet1.xml":
                s = data.decode()
                s = re.sub(r'<c r="F3"[^>]*>.*?</c>', '<c r="F3"><f>SUM(D2:D6)</f><v>2580</v></c>', s)
                s = s.replace("</sheetData>", "</sheetData>", 1)
                # inline string cell + string-typed formula result, appended as row 8
                extra = ('<row r="8"><c r="A8" t="inlineStr"><is><t>インライン文字列</t></is></c>'
                         '<c r="B8" t="str"><f>"a"&amp;"b"</f><v>ab</v></c>'
                         '<c r="C8" t="str"><v>str-typed value</v></c>'
                         '<c r="D8" t="b"><v>1</v></c><c r="E8" t="e"><v>#REF!</v></c></row>')
                errs = "".join(f'<c r="{col}9" t="e"><v>{e}</v></c>' for col, e in zip("ABCDEFG", [
                    "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"]))
                extra += f'<row r="9">{errs}</row>'
                s = s.replace("</sheetData>", extra + "</sheetData>", 1)
                data = s.encode()
            zout.writestr(n, data)
    os.remove(tmp)
    print("wrote", final)


def newer_errors():
    """newerrors.xlsx: error codes newer than the classic seven (written by Excel 365)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "スピル"
    for n, code in [("スピル", "#SPILL!"), ("カルク", "#CALC!"), ("フィールド", "#FIELD!"), ("取得中", "#GETTING_DATA")]:
        if n != "スピル":
            ws = wb.create_sheet(n)
        ws["A1"] = "before"
        ws["B1"] = "PLACEHOLDER"
        ws["C1"] = "after"
        ws["A2"] = code + " above"
    wb.create_sheet("正常")["A1"] = "fine"
    tmp = os.path.join(OUT, "newerrors.tmp.xlsx")
    wb.save(tmp)
    final = os.path.join(OUT, "newerrors.xlsx")
    codes = {1: "#SPILL!", 2: "#CALC!", 3: "#FIELD!", 4: "#GETTING_DATA"}
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(final, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = zin.read(n)
            m = re.match(r"xl/worksheets/sheet(\d)\.xml", n)
            if m and int(m.group(1)) in codes:
                x = data.decode()
                x = re.sub(r'<c r="B1"[^>]*>.*?</c>', f'<c r="B1" t="e"><v>{codes[int(m.group(1))]}</v></c>', x)
                data = x.encode()
            zout.writestr(n, data)
    os.remove(tmp)
    print("wrote", final)


main()
newer_errors()
